"""Append measured core-count sweep and confirmation without changing model predictions."""

from pathlib import Path
import json
import shutil
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import ScatterChart, Reference, Series

root = Path(__file__).resolve().parents[2]
base = root / "ttnn/ttnn/operations/examples/dram_l1_transaction_size"
sweep = json.loads((base / "core_sweep.json").read_text())
confirm = json.loads((base / "core_sweep_confirmation.json").read_text())
confirmed = {r["cores"]: r for r in confirm["results"]}
path = root / "Data movement perf calculator (1).xlsx"
backup = path.with_name(path.stem + ".before-core-sweep.xlsx")
if not backup.exists():
    shutil.copy2(path, backup)
w = openpyxl.load_workbook(path)
ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
caches = {}
with zipfile.ZipFile(path) as z:
    for i, ws in enumerate(w, 1):
        tree = ET.fromstring(z.read(f"xl/worksheets/sheet{i}.xml"))
        for c in tree.findall(".//s:sheetData/s:row/s:c", ns):
            if c.find("s:f", ns) is not None:
                v = c.find("s:v", ns)
                caches[(ws.title, c.attrib["r"])] = (c.attrib.get("t"), None if v is None else v.text)
if "Measured core sweep" in w:
    del w["Measured core sweep"]
s = w.create_sheet("Measured core sweep")
s.sheet_view.showGridLines = False
s.freeze_panes = "C11"
s.merge_cells("A1:J1")
s["A1"] = "Measured core-count sweep — preferred read NoC0"
s["A1"].fill = PatternFill("solid", fgColor="17365D")
s["A1"].font = Font(size=17, bold=True, color="FFFFFF")
notes = [
    f"{sweep['host']} | {sweep['arch']} | {sweep['clock_mhz']} MHz | {sweep['date']}",
    "BF16 [1,1,4096,1024], 8 MiB, 2048 B/transaction, one final barrier/core; row-major core prefixes.",
    f"All {len(sweep['results'])} feasible counts 7..110 tested (5 trials); 15 selected counts confirmed with 30 trials. Exact data and padding checks passed.",
    "Current runtime worker grid is 11 x 10 = 110. Counts 111..120 from the hardware model are unavailable in this session.",
    "Initial fastest sampled median: 88 cores / 21.847 us. Confirmation fastest sampled median: 44 cores / 21.875 us.",
    "44..99 selected counts form a near-tied plateau. 11 cores: 22.1495 us (confirmation), only 1.25% slower than 44 cores.",
    "Original model predicts 16.384 us for every feasible core count at 2048 B. Prediction inputs and calibration curves are unchanged.",
    "Times are device kernel duration, median and population std. Host copies excluded. Different counts select different prefixes of the fixed placement rule.",
]
for r, note in enumerate(notes, 2):
    s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    s.cell(r, 1, note).alignment = Alignment(wrap_text=True)
    s.row_dimensions[r].height = 30
labels = [
    "Cores",
    "L1 B/core",
    "Expected us (snapshot)",
    "Measured us (5 trials)",
    "Std us (5 trials)",
    "GB/s (5 trials)",
    "Confirmed us (30 trials)",
    "Std us (30 trials)",
    "GB/s (30 trials)",
    "Correct",
]
for c, label in enumerate(labels, 1):
    cell = s.cell(10, c, label)
    cell.fill = PatternFill("solid", fgColor="17365D")
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(wrap_text=True)
s.row_dimensions[10].height = 44
for r, item in enumerate(sweep["results"], 11):
    conf = confirmed.get(item["cores"])
    values = [
        item["cores"],
        item["allocated_l1_bytes_per_core"],
        item["expected_us"],
        item["median_us"],
        item["std_us"],
        item["gbps"],
        conf["median_us"] if conf else None,
        conf["std_us"] if conf else None,
        conf["gbps"] if conf else None,
        "PASS",
    ]
    for c, value in enumerate(values, 1):
        s.cell(r, c, value)
    for c in range(3, 10):
        s.cell(r, c).number_format = "0.000"
for c in range(1, 11):
    s.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 22
end = 10 + len(sweep["results"])
s.auto_filter.ref = f"A10:J{end}"
chart = ScatterChart()
chart.title = "Predicted vs measured duration by core count"
chart.x_axis.title = "Active cores"
chart.y_axis.title = "Microseconds"
chart.width = 27
chart.height = 14
for col in (3, 4, 7):
    series = Series(
        Reference(s, min_col=col, min_row=10, max_row=end),
        Reference(s, min_col=1, min_row=11, max_row=end),
        title_from_data=True,
    )
    chart.series.append(series)
s.add_chart(chart, "L10")
w.save(path)
with zipfile.ZipFile(path) as z:
    payload = {n: z.read(n) for n in z.namelist()}
for i, ws in enumerate(w, 1):
    name = f"xl/worksheets/sheet{i}.xml"
    tree = ET.fromstring(payload[name])
    changed = False
    for c in tree.findall(".//s:sheetData/s:row/s:c", ns):
        if (ws.title, c.attrib["r"]) not in caches:
            continue
        typ, val = caches[(ws.title, c.attrib["r"])]
        v = c.find("s:v", ns)
        if v is None:
            v = ET.SubElement(c, "{" + ns["s"] + "}v")
        v.text = val
        if typ:
            c.set("t", typ)
        else:
            c.attrib.pop("t", None)
        changed = True
    if changed:
        payload[name] = ET.tostring(tree, encoding="utf-8")
with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
    for name, data in payload.items():
        z.writestr(name, data)
check = openpyxl.load_workbook(path, data_only=True)
assert check["Model"]["B33"].value == 16.384
assert check["Measured core sweep"].max_row == 114
print("Added Measured core sweep; preserved original model and previous measurements.")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt

fig, (ax, bx) = plt.subplots(2, 1, figsize=(9.5, 7), layout="constrained", sharex=True)
x = [r["cores"] for r in sweep["results"]]
ax.plot(x, [r["expected_us"] for r in sweep["results"]], "--", label="Workbook prediction")
ax.plot(x, [r["median_us"] for r in sweep["results"]], ".-", linewidth=1, label="5-trial sweep: median")
ax.errorbar(
    [r["cores"] for r in confirm["results"]],
    [r["median_us"] for r in confirm["results"]],
    yerr=[r["std_us"] for r in confirm["results"]],
    fmt="o",
    ms=4,
    capsize=3,
    label="30-trial confirmation: median ± std",
)
ax.set_ylabel("Device duration (µs)")
ax.set_title("DRAM → L1 | BF16 [1,1,4096,1024] | 2048 B transactions | NoC0\nBH at 1350 MHz · bh-qbge-09 · 2026-10-08")
ax.legend()
ax.grid(alpha=0.2)
bx.plot(x, [r["gbps"] for r in sweep["results"]], ".-", linewidth=1)
bx.axhline(512, color="gray", linestyle="--", label="Configured theoretical peak: 512 GB/s")
bx.set_xlabel("Active cores (row-major placement)")
bx.set_ylabel("Aggregate bandwidth (GB/s)")
bx.grid(alpha=0.2)
bx.legend()
fig.savefig(base / "core_sweep.png", dpi=180)
