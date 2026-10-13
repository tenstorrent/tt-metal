"""Export current cached workbook shape and matching fixed-core predictions for a device run."""

from pathlib import Path
import json, hashlib
import openpyxl

root = Path(__file__).resolve().parents[2]
path = root / "Data movement perf calculator (1).xlsx"
w = openpyxl.load_workbook(path, data_only=True)
m = w["Model"]
cores = int(m["B29"].value)
assert m["B8"].value == 2, "Example currently supports BF16 storage"
case = dict(
    shape=[int(m.cell(r, 2).value) for r in (4, 5, 6, 7)],
    cores=cores,
    bytes_per_element=2,
    model_clock_ghz=w["Hardware"]["B4"].value,
    workbook_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
    hardware={str(w["Hardware"].cell(r, 1).value): w["Hardware"].cell(r, 2).value for r in range(4, 20)},
    predictions=[],
    core_predictions=[],
)
for row in w["Search"].iter_rows(min_row=4, values_only=True):
    if row[20] == 1 and row[1] == 2048:
        case["core_predictions"].append(
            dict(cores=row[0], transaction_bytes=row[1], expected_us=row[19], max_shard_bytes=row[4])
        )
    if row[0] == cores and row[20] == 1:
        case["predictions"].append(
            dict(
                transaction_bytes=row[1],
                expected_us=row[19],
                core_service_us=row[16],
                dram_floor_us=row[17],
                max_shard_bytes=row[4],
                transactions_max_core=row[5],
            )
        )
assert len(case["predictions"]) > 0
out = root / "tools/data_movement_perf/benchmark_case.json"
out.write_text(json.dumps(case, indent=2) + "\n")
print(out)
print(json.dumps(case, indent=2))
