"""Set the model's aggregate DRAM efficiency and recalculate Excel formulas locally."""

from pathlib import Path
import functools, json, shutil, zipfile, xml.etree.ElementTree as ET
import numpy as np
import openpyxl
import formulas
from openpyxl.styles import Font, PatternFill, Alignment

root = Path(__file__).resolve().parents[2]
p = root / "Data movement perf calculator (1).xlsx"
backup = p.with_name(p.stem + ".before-efficiency.xlsx")
if not backup.exists():
    shutil.copy2(p, backup)
w = openpyxl.load_workbook(p)
old = openpyxl.load_workbook(p, data_only=True)
w["Hardware"]["B15"] = 0.70
w["Hardware"][
    "D15"
] = "70% aggregate DRAM ceiling assumption = 358.4 GB/s. Conservative relative to this workload's measured ~75% plateau; does not model core-count scaling."
active = {"Hardware", "Model", "Search", "Core summary", "Shard plan"}
names = {s.upper(): s for s in w.sheetnames}


@functools.lru_cache(None)
def value(sn, addr):
    c = w[sn][addr]
    if c.data_type != "f":
        return c.value if c.value is not None else ""
    if sn not in active:
        v = old[sn][addr].value
        return v if v is not None else ""
    fn = formulas.Parser().ast(c.value)[1].compile()
    args = []
    for ref in fn.inputs:
        if "!" in ref:
            sheet, a = ref.rsplit("!", 1)
            sheet = names[sheet.strip("'")]
        else:
            sheet, a = sn, ref
        cells = w[sheet][a.replace("$", "")]
        if isinstance(cells, tuple):
            arg = np.array([[value(sheet, x.coordinate) for x in row] for row in cells], dtype=object)
        else:
            arg = value(sheet, cells.coordinate)
        args.append(arg)
    v = np.asarray(fn(*args)).item()
    if isinstance(v, np.generic):
        v = v.item()
    if not isinstance(v, (str, float, int, bool)):
        raise RuntimeError((sn, addr, v))
    if isinstance(v, str) and v.startswith(("#REF!", "#VALUE!", "#DIV/0!", "#NUM!")):
        raise RuntimeError((sn, addr, v))
    return v


# Evaluate real formulas, including updated rank/recommendation and shard assignments.
caches = {}
for sn in w.sheetnames:
    for row in w[sn]:
        for c in row:
            if c.data_type == "f":
                caches[(sn, c.coordinate)] = value(sn, c.coordinate)
print("Recalculated model formulas", flush=True)
base = root / "ttnn/ttnn/operations/examples/dram_l1_transaction_size"
measured = json.loads((base / "core_sweep_confirmation.json").read_text())
if "DRAM efficiency comparison" in w:
    del w["DRAM efficiency comparison"]
s = w.create_sheet("DRAM efficiency comparison")
s.freeze_panes = "B9"
s.sheet_view.showGridLines = False
s.merge_cells("A1:H1")
s["A1"] = "Aggregate DRAM efficiency — predictions versus measured NoC0 reads"
s["A1"].font = Font(size=16, bold=True, color="FFFFFF")
s["A1"].fill = PatternFill("solid", fgColor="17365D")
notes = [
    "Active model: Hardware!B15 = 70%; chip bandwidth ceiling 358.4 GB/s. 75% sensitivity column uses 384 GB/s.",
    "Shape BF16 [1,1,4096,1024], 8 MiB, 2048 B transactions, one barrier/core, preferred read NoC0; BH 1350 MHz.",
    "Measurements: 30-trial medians from the fixed row-major placement sweep. No new hardware measurements were made.",
    "70% and 75% below are snapshot predictions for this shape. The main Model/Search sheets recalculate using editable Hardware!B15.",
    "This changes the aggregate ceiling only. It cannot reproduce the measured dependence on core count; seven-core estimates remain optimistic.",
    "75% is an in-sample fit to the observed plateau, not independent validation. Earlier comparison sheets retain their original predictions.",
]
for r, note in enumerate(notes, 2):
    s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    s.cell(r, 1, note).alignment = Alignment(wrap_text=True)
    s.row_dimensions[r].height = 31
heads = [
    "Cores",
    "Measured us",
    "Original 100% us",
    "70% prediction us",
    "70% error %",
    "75% prediction us",
    "75% error %",
    "Measured GB/s",
]
for col, h in enumerate(heads, 1):
    c = s.cell(8, col, h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="17365D")
    c.alignment = Alignment(wrap_text=True)
s.row_dimensions[8].height = 38
results = []
for ri, r in enumerate(measured["results"], 9):
    # Match actual transaction/core candidate. Change shared/local bank floors only for 75% sensitivity.
    searchrow = 4 + (r["cores"] - 1) * 9 + 5
    assert w["Search"].cell(searchrow, 2).value == 2048
    t70 = value("Search", f"T{searchrow}")
    t75 = max(
        value("Search", f"Q{searchrow}"),
        value("Search", f"R{searchrow}") * 0.70 / 0.75,
        value("Search", f"S{searchrow}") * 0.70 / 0.75,
    ) + value("Model", "B19")
    e70 = (t70 / r["median_us"] - 1) * 100
    e75 = (t75 / r["median_us"] - 1) * 100
    vals = [r["cores"], r["median_us"], r["expected_us"], t70, e70, t75, e75, r["gbps"]]
    for col, v in enumerate(vals, 1):
        s.cell(ri, col, v)
    for col in range(2, 9):
        s.cell(ri, col).number_format = "0.000"
    results.append(
        dict(
            cores=r["cores"],
            measured_us=r["median_us"],
            predicted_70_us=t70,
            error_70_percent=e70,
            predicted_75_us=t75,
            error_75_percent=e75,
        )
    )
for col in range(1, 9):
    s.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 24
s.auto_filter.ref = f"A8:H{8+len(results)}"
# Preserve all prior cached formula results and write new values for the active model.
w.save(p)
ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
with zipfile.ZipFile(p) as z:
    payload = {n: z.read(n) for n in z.namelist()}
for i, ws in enumerate(w, 1):
    name = f"xl/worksheets/sheet{i}.xml"
    tree = ET.fromstring(payload[name])
    changed = False
    for c in tree.findall(".//s:sheetData/s:row/s:c", ns):
        key = (ws.title, c.attrib["r"])
        if key not in caches:
            continue
        val = caches[key]
        v = c.find("s:v", ns)
        if v is None:
            v = ET.SubElement(c, "{" + ns["s"] + "}v")
        if isinstance(val, str):
            c.set("t", "str")
            v.text = val
        else:
            c.attrib.pop("t", None)
            v.text = str(int(val) if isinstance(val, bool) else val)
        changed = True
    if changed:
        payload[name] = ET.tostring(tree, encoding="utf-8")
with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
    for name, data in payload.items():
        z.writestr(name, data)
check = openpyxl.load_workbook(p, data_only=True)
assert abs(check["Model"]["B33"].value - 8388608 / (358.4 * 1000)) < 1e-8
assert check["Model"]["B29"].value == 7
assert check["Measured comparison"]["E15"].value == 27.144
out = dict(
    efficiency=0.70,
    bandwidth_cap_gbps=358.4,
    recommended_cores=value("Model", "B29"),
    recommended_transaction=value("Model", "B30"),
    recommended_us=value("Model", "B33"),
    results=results,
)
(root / "tools/data_movement_perf/efficiency_comparison.json").write_text(json.dumps(out, indent=2) + "\n")
print(json.dumps(out, indent=2))
