"""Append measured results without changing model inputs, formulas, or cached predictions."""

from pathlib import Path
import json, shutil, zipfile, xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

root = Path(__file__).resolve().parents[2]
path = root / "Data movement perf calculator (1).xlsx"
example = root / "ttnn/ttnn/operations/examples/dram_l1_transaction_size"
report = json.loads((example / "report.json").read_text())
steady = json.loads((example / "steady_state.json").read_text())
backup = path.with_name(path.stem + ".before-benchmark.xlsx")
if not backup.exists():
    shutil.copy2(path, backup)
w = openpyxl.load_workbook(path)
# openpyxl clears formula caches when saving: restore the exact caches of existing sheets.
with zipfile.ZipFile(path) as z:
    old = {n: z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")}
ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
cached = {}
for i, ws in enumerate(w, 1):
    raw = old.get(f"xl/worksheets/sheet{i}.xml")
    if raw is None:
        continue
    for c in ET.fromstring(raw).findall(".//s:sheetData/s:row/s:c", ns):
        if c.find("s:f", ns) is not None:
            v = c.find("s:v", ns)
            cached[(ws.title, c.attrib["r"])] = (c.attrib.get("t"), v.text if v is not None else None)
if "Measured comparison" in w:
    if "Measured alternating NoCs" not in w:
        w["Measured comparison"].title = "Measured alternating NoCs"
        cached = {
            ("Measured alternating NoCs" if sn == "Measured comparison" else sn, cell): value
            for (sn, cell), value in cached.items()
        }
    else:
        del w["Measured comparison"]
s = w.create_sheet("Measured comparison")
s.sheet_view.showGridLines = False
s.freeze_panes = "C10"
s.merge_cells("A1:J1")
s["A1"] = "Measured DRAM → L1: transaction-size sweep"
s["A1"].fill = PatternFill("solid", fgColor="17365D")
s["A1"].font = Font(size=17, bold=True, color="FFFFFF")
metadata = [
    ("Device", f"{report['host']} | {report['arch']} | {report['clock_mhz']} MHz"),
    ("Workload", f"{report['shape']} BF16 | 8 MiB | {report['cores']} cores | one final barrier/core"),
    (
        "Placement",
        f"Logical cores {report['logical_cores']}; {report.get('noc_policy', 'alternating 0/1')}; NCRISC reader",
    ),
    (
        "Measurement",
        f"DEVICE KERNEL DURATION [ns]; {report['trials']} trials; median and population std; warmup discarded",
    ),
    (
        "Correctness",
        "Exact BF16 comparison of all rows; untouched shard padding checked; host transfers excluded from timing",
    ),
    ("Recorded", f"{report['date']} | model snapshot SHA256 {report['model']['workbook_sha256']}"),
]
for r, (label, value) in enumerate(metadata, 2):
    s.cell(r, 1, label)
    s.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
    s.cell(r, 2, value).alignment = Alignment(wrap_text=True)
    s.row_dimensions[r].height = 30
labels = [
    "Cores",
    "Transaction B",
    "Transactions / largest core",
    "Expected us (snapshot)",
    "Measured us (1 iteration)",
    "Std us",
    "Measured / expected",
    "Measured GB/s",
    "Steady us (16 iterations)",
    "Correct",
]
for i, label in enumerate(labels, 1):
    c = s.cell(9, i, label)
    c.fill = PatternFill("solid", fgColor="17365D")
    c.font = Font(bold=True, color="FFFFFF")
    c.alignment = Alignment(wrap_text=True)
s.row_dimensions[9].height = 45
newcache = {}
for r, (v, t) in enumerate(zip(report["results"], steady["results"]), 10):
    assert v["transaction_bytes"] == t["transaction_bytes"]
    values = [
        report["cores"],
        v["transaction_bytes"],
        v["transactions_max_core"],
        v["expected_us"],
        v["median_us"],
        v["std_us"],
        f"=E{r}/D{r}",
        f'={report["row_bytes"]*report["shape"][0]*report["shape"][1]*report["shape"][2]}/(E{r}*1000)',
        t["median_us"],
        "PASS",
    ]
    for c, value in enumerate(values, 1):
        s.cell(r, c, value)
    for c in range(4, 10):
        s.cell(r, c).number_format = "0.000"
    newcache[(s.title, f"G{r}")] = (None, str(v["ratio"]))
    newcache[(s.title, f"H{r}")] = (None, str(v["gbps"]))
notes = [
    "Expected values are frozen from the model used at measurement time. Editing Model does not update these recorded predictions.",
    f"2048 B payload: {report['results'][-1]['median_us']:.3f} us measured vs {report['results'][-1]['expected_us']:.3f} us predicted.",
    f"16-iteration control: {steady['results'][-1]['median_us']:.3f} us/transfer; launch overhead alone does not explain the model gap.",
    f"Measured throughput is {report['results'][-1]['gbps']:.2f} GB/s. Original model predictions use unchanged hardware limits and are not NoC-specific.",
    "The measured single-bank/all-bank curves do not capture this exact kernel, per-request issue cost, core placement, or multicore contention.",
    "L1 allocation is 1200128 B/core (uniform 586-row capacity). Six cores read 585 rows and retain one unused padding row. No extra staging buffer.",
    "Reproduction: see ttnn/ttnn/operations/examples/dram_l1_transaction_size/README.md; original predictions and Hardware inputs are unchanged.",
]
for r, note in enumerate(notes, 18):
    s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
    s.cell(r, 1, note).alignment = Alignment(wrap_text=True)
    s.row_dimensions[r].height = 30
for col in range(1, 11):
    s.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 21
s.auto_filter.ref = "A9:J15"
chart = BarChart()
chart.title = "Predicted vs measured device duration"
chart.y_axis.title = "Microseconds"
chart.x_axis.title = "Transaction bytes"
chart.add_data(Reference(s, min_col=4, max_col=5, min_row=9, max_row=15), titles_from_data=True)
chart.set_categories(Reference(s, min_col=2, min_row=10, max_row=15))
chart.width = 25
chart.height = 12
s.add_chart(chart, "A27")
w.save(path)
cached.update(newcache)
with zipfile.ZipFile(path) as z:
    payload = {n: z.read(n) for n in z.namelist()}
for i, ws in enumerate(w, 1):
    name = f"xl/worksheets/sheet{i}.xml"
    tree = ET.fromstring(payload[name])
    changed = False
    for c in tree.findall(".//s:sheetData/s:row/s:c", ns):
        key = (ws.title, c.attrib["r"])
        if key not in cached:
            continue
        typ, val = cached[key]
        v = c.find("s:v", ns)
        if v is None:
            v = ET.SubElement(c, "{" + ns["s"] + "}v")
        v.text = val
        if typ:
            c.set("t", typ)
        else:
            c.attrib.pop("t", None)
        changed = True
    if changed:
        payload[name] = ET.tostring(tree, encoding="utf-8")
with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
    for name, data in payload.items():
        z.writestr(name, data)
check = openpyxl.load_workbook(path, data_only=True)
assert check["Model"]["B33"].value == 16.384
assert abs(check[s.title]["G15"].value - report["results"][-1]["ratio"]) < 1e-10
print("Added Measured comparison; model inputs and formula caches preserved.")
# Shareable figure, with fixed units and full provenance in the accompanying report.
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt

xs = [r["transaction_bytes"] for r in report["results"]]
fig, ax = plt.subplots(figsize=(8.6, 4.8), layout="constrained")
ax.plot(xs, [r["expected_us"] for r in report["results"]], "o--", label="Workbook prediction")
ax.errorbar(
    xs,
    [r["median_us"] for r in report["results"]],
    yerr=[r["std_us"] for r in report["results"]],
    fmt="o-",
    capsize=3,
    label="Measured: 1 transfer/launch",
)
ax.plot(xs, [r["median_us"] for r in steady["results"]], "s:", label="Measured: 16 transfers/launch, normalized")
ax.set_xscale("log", base=2)
ax.set_yscale("log")
ax.set_xticks(xs, [str(v) for v in xs])
ax.set_xlabel("Transaction payload (bytes)")
ax.set_ylabel("Duration per transfer (µs)")
ax.set_title("DRAM → L1 | BF16 [1, 1, 4096, 1024] | 7 cores\nBH, 1350 MHz · bh-qbge-09 · 5 trials · 2026-10-08")
ax.grid(True, alpha=0.2)
ax.legend()
fig.savefig(example / "comparison.png", dpi=180)
