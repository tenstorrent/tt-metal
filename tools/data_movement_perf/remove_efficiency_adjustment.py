"""Remove unsupported timing adjustments and recalculate the workbook."""

from pathlib import Path
import functools, json, shutil, zipfile, xml.etree.ElementTree as ET
import numpy as np
import openpyxl
import formulas
from openpyxl.styles import Font, PatternFill, Alignment

root = Path(__file__).resolve().parents[2]
p = root / "Data movement perf calculator (1).xlsx"
backup = p.with_name(p.stem + ".before-removing-efficiency.xlsx")
if not backup.exists():
    shutil.copy2(p, backup)
w = openpyxl.load_workbook(p)
old = openpyxl.load_workbook(p, data_only=True)
w["Hardware"]["B15"] = 1
w["Hardware"]["A15"] = "Legacy efficiency (unused)"
w["Hardware"]["D15"] = "70% assumption removed. Active timing formulas use unscaled theoretical DRAM bandwidth."
w["Model"]["B24"] = "Unscaled DRAM floor + empirical read envelope; incomplete congestion model"
pipeline = w["Read pipeline"]
pipeline["A2"] = "Component startup enabled: 0 omitted, 1 explicit inputs"
pipeline["B2"] = 0
pipeline["C2"] = "No separate startup addition by default; missing overhead is unmodeled, not known to be zero."
pipeline["B12"] = "=IF(B2=1,B11,0)"
pipeline["A19"] = "Aggregate efficiency adjustment"
pipeline["B19"] = "Removed from active formulas"
pipeline["C19"] = "The 70% cap is not an explained hardware limit."
pipeline["B20"] = "Unscaled theoretical DRAM floor + empirical read envelope"
pipeline["C20"] = "Measured curves retained without subtracting or adding a latency proxy."
pipeline["B21"] = "MAX(empirical duration, physical service bounds)"
pipeline["C21"] = "Optional explicit component startup applies only to physical bounds; off by default."
pipeline["C22"] = "Shared-link congestion is not modeled. No efficiency factor substitutes for it."
pipeline["C23"] = "Composite latency retained for reference; not added separately to transfer time."
search = w["Search"]
search["AA3"] = "Empirical read duration us"
for r in range(4, 1084):
    for col in ["R", "S"]:
        search[f"{col}{r}"] = search[f"{col}{r}"].value.replace("*Hardware!$B$15", "")
    search[f"AA{r}"] = f"=M{r}"
    search[f"AF{r}"] = "='Read pipeline'!$B$12"
    search[f"O{r}"] = f"=MAX(AA{r},AF{r}+MAX(AB{r},AC{r},AD{r},AE{r}))"
    search[f"T{r}"] = f"=MAX(AA{r},AF{r}+MAX(AB{r},AC{r},AD{r},AE{r},R{r},S{r}))+P{r}+Model!$B$19"
for row in w["Guide"]:
    if row[0].value == "Read pipeline":
        row[
            1
        ].value = "Active estimate: MAX(empirical read duration, unscaled physical service bounds). No 70% cap or separate composite-latency addition. Historical sheets retain earlier assumptions."
w["Guide"].append(
    [
        "Unknown costs",
        "Issue/window constraints and explicit startup are disabled pending evidence. Zero means omitted, not physically absent. Empirical curves and placement assumptions limit generalization.",
    ]
)
active = {"Hardware", "Model", "Search", "Core summary", "Shard plan", "Read pipeline", "Pipeline comparison"}
names = {s.upper(): s for s in w.sheetnames}


@functools.lru_cache(None)
def value(sn, addr):
    c = w[sn][addr]
    if c.data_type != "f":
        return c.value if c.value is not None else ""
    if sn not in active:
        v = old[sn][addr].value
        return v if v is not None else ""
    fn = formulas.Parser().ast(c.value)[1].compile()
    args = []
    for ref in fn.inputs:
        if "!" in ref:
            sheet, a = ref.rsplit("!", 1)
            sheet = names[sheet.strip("'")]
        else:
            sheet, a = sn, ref
        cells = w[sheet][a.replace("$", "")]
        if isinstance(cells, tuple):
            arg = np.array([[value(sheet, x.coordinate) for x in row] for row in cells], dtype=object)
        else:
            arg = value(sheet, cells.coordinate)
        args.append(arg)
    v = np.asarray(fn(*args)).item()
    if isinstance(v, np.generic):
        v = v.item()
    if not isinstance(v, (str, float, int, bool)):
        raise RuntimeError((sn, addr, v))
    if isinstance(v, str) and v.startswith(("#REF!", "#VALUE!", "#DIV/0!", "#NUM!")):
        raise RuntimeError((sn, addr, v))
    return v


base = root / "ttnn/ttnn/operations/examples/dram_l1_transaction_size"
measured = json.loads((base / "core_sweep_confirmation.json").read_text())
if "Pipeline comparison" in w:
    del w["Pipeline comparison"]
s = w.create_sheet("Pipeline comparison")
s.append(
    [
        "Cores",
        "Measured us (existing 30 trials)",
        "Previous estimate (snapshot) us",
        "Unscaled estimate us",
        "Error %",
        "Measured GB/s",
    ]
)
results = []
for ri, r in enumerate(measured["results"], 2):
    sr = 4 + (r["cores"] - 1) * 9 + 5
    pred = value("Search", f"T{sr}")
    s.append(
        [
            r["cores"],
            r["median_us"],
            old["Search"][f"T{sr}"].value,
            f"=Search!T{sr}",
            f"=(D{ri}/B{ri}-1)*100",
            r["gbps"],
        ]
    )
    results.append(
        dict(
            cores=r["cores"],
            measured_us=r["median_us"],
            predicted_us=pred,
            error_percent=(pred / r["median_us"] - 1) * 100,
        )
    )
for col in "ABCDEF":
    s.column_dimensions[col].width = 28
names = {s.upper(): s for s in w.sheetnames}
# Evaluate real formulas, including updated rank/recommendation and shard assignments.
caches = {}
for sn in w.sheetnames:
    for row in w[sn]:
        for c in row:
            if c.data_type == "f":
                caches[(sn, c.coordinate)] = value(sn, c.coordinate)
# Preserve all prior cached formula results and write new values for the active model.
w.save(p)
ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
with zipfile.ZipFile(p) as z:
    payload = {n: z.read(n) for n in z.namelist()}
for i, ws in enumerate(w, 1):
    name = f"xl/worksheets/sheet{i}.xml"
    tree = ET.fromstring(payload[name])
    changed = False
    for c in tree.findall(".//s:sheetData/s:row/s:c", ns):
        key = (ws.title, c.attrib["r"])
        if key not in caches:
            continue
        val = caches[key]
        v = c.find("s:v", ns)
        if v is None:
            v = ET.SubElement(c, "{" + ns["s"] + "}v")
        if isinstance(val, str):
            c.set("t", "str")
            v.text = val
        else:
            c.attrib.pop("t", None)
            v.text = str(int(val) if isinstance(val, bool) else val)
        changed = True
    if changed:
        payload[name] = ET.tostring(tree, encoding="utf-8")
with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
    for name, data in payload.items():
        z.writestr(name, data)
check = openpyxl.load_workbook(p, data_only=True)
assert check["Measured comparison"]["E15"].value == 27.144
out = dict(
    recommended_cores=value("Model", "B29"),
    transaction_bytes=value("Model", "B30"),
    predicted_us=value("Model", "B33"),
    results=results,
)
(root / "tools/data_movement_perf/pipeline_comparison.json").write_text(json.dumps(out, indent=2) + "\n")
print(json.dumps(out, indent=2))
