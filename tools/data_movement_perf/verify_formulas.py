"""Independently evaluate representative Excel formulas against reference caches."""
from pathlib import Path
import math
import numpy as np
import openpyxl
import formulas

p = Path(__file__).resolve().parents[2] / "Data movement perf calculator (1).xlsx"
f = openpyxl.load_workbook(p, data_only=False)
v = openpyxl.load_workbook(p, data_only=True)
names = {s.upper(): s for s in v.sheetnames}
count = 0
for sheet in ["Hardware", "Model", "Search", "Core summary", "Shard plan", "Read pipeline", "Pipeline comparison"]:
    for row in f[sheet]:
        for cell in row:
            if cell.data_type != "f":
                continue
            fn = formulas.Parser().ast(cell.value)[1].compile()
            args = []
            for ref in fn.inputs:
                if "!" in ref:
                    sn, addr = ref.rsplit("!", 1)
                    sn = names[sn.strip("'")]
                else:
                    sn, addr = sheet, ref
                data = v[sn][addr.replace("$", "")]
                args.append(
                    np.array([[c.value for c in r] for r in data], dtype=object)
                    if isinstance(data, tuple)
                    else (data.value if data.value is not None else "")
                )
            got = np.asarray(fn(*args)).item()
            want = v[sheet][cell.coordinate].value
            if isinstance(want, (int, float)):
                assert isinstance(got, (int, float, np.number)) and math.isclose(
                    float(got), want, rel_tol=1e-10, abs_tol=1e-9
                ), (sheet, cell.coordinate, got, want)
            else:
                assert got == (want or ""), (sheet, cell.coordinate, got, want)
            count += 1
print(f"PASS: {count} Excel formulas independently evaluated against reference results.")
